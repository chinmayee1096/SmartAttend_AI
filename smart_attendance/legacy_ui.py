import streamlit as st
import cv2
import os
import sys
import numpy as np
import pandas as pd
from pathlib import Path
from PIL import Image
from datetime import datetime, timedelta
import time
import traceback
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
import io
from face_pipeline import prepare_face, valid_sample, detect_faces

# Try importing optional libraries
try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import inch
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Check if running with streamlit run
try:
    from streamlit.runtime.scriptrunner import get_script_run_ctx
    if not get_script_run_ctx():
        print("\n\n" + "="*60)
        print("ERROR: You are running this script incorrectly!")
        print("Please do NOT run 'python streamlit_app.py'.")
        print("Instead, run: streamlit run streamlit_app.py")
        print("="*60 + "\n\n")
except ImportError:
    pass

# -----------------------
# Configuration & Helpers
# -----------------------
BASE_DIR = Path(__file__).resolve().parents[1] / "dataset" / "college_faces"
BASE_DIR.mkdir(parents=True, exist_ok=True)
DEPARTMENTS = ["AIML", "CSE", "Civil", "ECE", "MECH"]
SECTIONS = ["A", "B", "C"]

def create_lbph_recognizer():
    if hasattr(cv2, "face") and hasattr(cv2.face, "LBPHFaceRecognizer_create"):
        return cv2.face.LBPHFaceRecognizer_create()
    try:
        return cv2.face.LBPHFaceRecognizer_create()
    except Exception:
        st.error("LBPH recognizer not available. Please install opencv-contrib-python.")
        return None

def get_schedule():
    def t(h, m): return datetime.strptime(f"{h}:{m}", "%H:%M").time()
    return [
        ("Period1",  t(8,30),  t(9,30)),
        ("Period2",  t(9,30),  t(10,30)),
        ("Break1",   t(10,30), t(11,0)),
        ("Period3",  t(11,0),  t(12,0)),
        ("Period4",  t(12,0),  t(13,0)),
        ("Break2",   t(13,0),  t(13,40)),
        ("Period5",  t(13,40), t(14,40)),
        ("Period6",  t(14,40), t(15,40)),
        ("Period7",  t(18,0),  t(19,0)),
        ("Period8",  t(19,0),  t(20,0)),
        ("Period9",  t(20,0),  t(21,0)),
        ("Period10", t(21,0),  t(22,0)),
        ("Period11", t(22,0),  t(23,0)),
        ("Period12", t(23,0),  t(23,59)),
    ]

def get_current_period():
    now = datetime.now().time()
    for name, start, end in get_schedule():
        if start <= now <= end:
            return name, start, end
    return None, None, None

def get_attendance_status(sign_in_time_str, period_start):
    """
    Determines if student is Present or Late.
    Late = arrived more than 10 minutes after period started.
    """
    try:
        sign_in_time = datetime.strptime(sign_in_time_str, "%Y-%m-%d %H:%M:%S").time()
        period_start_dt = datetime.combine(datetime.today(), period_start)
        sign_in_dt = datetime.combine(datetime.today(), sign_in_time)
        diff_minutes = (sign_in_dt - period_start_dt).total_seconds() / 60
        if diff_minutes > 10:
            return "Late"
        return "Present"
    except Exception:
        return "Present"

# -----------------------
# Page: Training
# -----------------------
def open_camera():
    from smart_attendance.services.camera_service import manager, CameraError
    import uuid
    owner = uuid.uuid4().hex
    try:
        manager.start(owner, int(st.session_state.get("camera_index", 0)))
    except CameraError as exc:
        st.error(str(exc))
        return None
    class Lease:
        def read(self):
            try:
                for _ in range(30):
                    frame = manager.read(owner)
                    if frame is not None:
                        return True, frame
                    time.sleep(.05)
            except CameraError:
                return False, None
            return False, None
        def release(self):
            manager.close(owner)
    return Lease()


def get_face_detector():
    detector = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
    if detector.empty():
        raise RuntimeError("Face detector missing. Install the packages in requirements.txt and restart the app.")
    return detector


def page_training():
    st.header("🎓 Student Face Training")
    st.info("Capture one student at a time in even light. Keep eyes visible, face the camera, and vary your angle slightly. Blurry or poorly detected samples are skipped.")

    col1, col2 = st.columns(2)
    with col1:
        roll = st.text_input("Roll Number (numbers only)")
        name = st.text_input("Name")
    with col2:
        dept = st.selectbox("Department", DEPARTMENTS)
        section = st.selectbox("Section", SECTIONS)

    if "training_active" not in st.session_state:
        st.session_state.training_active = False

    start_btn = st.button("▶ Start Capture", disabled=st.session_state.training_active)
    stop_btn  = st.button("⏹ Stop Capture",  disabled=not st.session_state.training_active)

    if start_btn:
        if not roll or not name:
            st.error("Please enter Roll Number and Name.")
        else:
            try:
                int(roll)
            except ValueError:
                st.error("Roll Number must be a plain integer (e.g. 101, 12345).")
                return
            st.session_state.training_active   = True
            st.session_state.training_roll     = roll
            st.session_state.training_name     = name
            st.session_state.training_dept     = dept
            st.session_state.training_section  = section
            st.rerun()

    if stop_btn:
        st.session_state.training_active = False
        st.rerun()

    video_placeholder = st.empty()
    progress_bar      = st.progress(0)
    status_text       = st.empty()

    if st.session_state.training_active:
        roll    = st.session_state.training_roll
        dept    = st.session_state.training_dept
        section = st.session_state.training_section
        name    = st.session_state.training_name

        faces_dir = BASE_DIR / dept / section / "faces" / roll
        faces_dir.mkdir(parents=True, exist_ok=True)

        # Update students.csv
        dept_dir = BASE_DIR / dept / section
        dept_dir.mkdir(parents=True, exist_ok=True)
        csv_path = dept_dir / "students.csv"

        if csv_path.exists():
            df = pd.read_csv(csv_path, dtype=str)
        else:
            df = pd.DataFrame(columns=["roll_number", "name", "branch", "section", "email"])

        if not df["roll_number"].astype(str).eq(str(roll)).any():
            new_row = {"roll_number": roll, "name": name, "branch": dept, "section": section, "email": ""}
            df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
            df.to_csv(csv_path, index=False)

        cap          = open_camera()
        if cap is None:
            st.session_state.training_active = False
            return
        face_cascade = get_face_detector()

        sample_count    = 0
        required_samples = 50

        try:
            while st.session_state.training_active and sample_count < required_samples:
                ret, frame = cap.read()
                if not ret:
                    st.session_state.training_active = False
                    st.error("Camera stopped delivering frames. Restart capture.")
                    break
    
                gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
                faces = detect_faces(gray, face_cascade)
    
                for (x, y, w, h) in (faces if len(faces) == 1 else []):
                    cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 255, 0), 2)
                    face_img = gray[y:y+h, x:x+w]
                    try:
                        face_img = cv2.resize(face_img, (200, 200))
                        if cv2.imwrite(str(faces_dir / f"face_{sample_count}.jpg"), face_img):
                            sample_count += 1
                    except Exception:
                        pass
    
                progress_bar.progress(min(sample_count / required_samples, 1.0))
                status_text.text(f"Captured: {sample_count} / {required_samples}")
    
                frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                video_placeholder.image(frame_rgb, channels="RGB")
    
                if sample_count >= required_samples:
                    st.session_state.training_active = False
                    cap.release()
                    st.success(f"✅ Capture completed for {name} ({roll})! Now click 'Train Models'.")
                    st.rerun()
        finally:
            cap.release()

    st.markdown("---")
    if st.button("🤖 Train Models (All Departments)"):
        train_models()

    # ---- Email management ----
    st.markdown("---")
    st.subheader("📧 Manage Student Emails")
    st.info("Add parent/student email addresses here to enable absence alerts.")

    em_dept    = st.selectbox("Department", DEPARTMENTS, key="em_dept")
    em_section = st.selectbox("Section",    SECTIONS,    key="em_section")

    csv_path = BASE_DIR / em_dept / em_section / "students.csv"
    if csv_path.exists():
        df_em = pd.read_csv(csv_path, dtype=str)
        if "email" not in df_em.columns:
            df_em["email"] = ""

        edited_df = st.data_editor(
            df_em[["roll_number", "name", "email"]],
            width="stretch",
            key="email_editor"
        )

        if st.button("💾 Save Emails"):
            df_em["email"] = edited_df["email"].values
            df_em.to_csv(csv_path, index=False)
            st.success("Emails saved!")
    else:
        st.warning("No students registered for this department/section yet.")


def train_models():
    from train_faces import train_all
    with st.spinner("Checking samples and training models..."):
        report = train_all(BASE_DIR)
    load_resources.clear()
    st.dataframe(pd.DataFrame(report), width="stretch")
    st.success("Training finished. See accepted samples and validation results above.")


# -----------------------
# Page: Attendance
# -----------------------
def page_attendance():
    st.header("📋 Daily Attendance")
    st.checkbox("Preview only (do not record attendance)", key="preview_only")

    if "attendance_active" not in st.session_state:
        st.session_state.attendance_active = False

    col1, col2, col3 = st.columns(3)
    with col1:
        start_btn = st.button("▶ Start Recognition", disabled=st.session_state.attendance_active)
    with col2:
        stop_btn  = st.button("⏹ Stop Recognition",  disabled=not st.session_state.attendance_active)
    with col3:
        if st.button("🔄 Reload Data"):
            st.cache_resource.clear()
            st.success("Reloaded!")

    pname, pstart, pend = get_current_period()
    if pname:
        if "Break" in pname:
            st.warning(f"☕ Currently: **{pname}** ({pstart.strftime('%H:%M')} - {pend.strftime('%H:%M')}) — No attendance during breaks.")
        else:
            st.info(f"🕐 Current Period: **{pname}** ({pstart.strftime('%H:%M')} - {pend.strftime('%H:%M')})")
    else:
        st.warning("⚠️ No active period currently. Attendance will not be marked.")

    video_placeholder = st.empty()
    log_placeholder   = st.empty()

    if stop_btn:
        st.session_state.attendance_active = False
        st.rerun()

    if start_btn:
        st.session_state.attendance_active = True
        st.rerun()

    if st.session_state.attendance_active:
        run_attendance_loop(video_placeholder, log_placeholder)

    # ---- Reports Section ----
    st.markdown("---")
    st.subheader("📊 Attendance Reports")

    r_dept = st.selectbox("Select Department", DEPARTMENTS, key="rep_dept")
    date_str = st.date_input("Select Date", datetime.now()).strftime("%Y-%m-%d")

    report_path = BASE_DIR / "reports" / f"{r_dept}_{date_str}_attendance.csv"

    if report_path.exists():
        df_rep = pd.read_csv(report_path, dtype={"roll_number": str})
        
        # Color-code the status column
        def highlight_status(val):
            if val == "Present":
                return "background-color: #d4edda; color: #155724"
            elif val == "Late":
                return "background-color: #fff3cd; color: #856404"
            return ""

        styled = df_rep.style.map(highlight_status, subset=["status"])
        st.dataframe(styled, width="stretch")

        # Summary stats
        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("Total Records", len(df_rep))
        with col2:
            present_count = len(df_rep[df_rep["status"] == "Present"])
            st.metric("Present", present_count)
        with col3:
            late_count = len(df_rep[df_rep["status"] == "Late"])
            st.metric("Late", late_count)

        st.markdown("---")
        st.subheader("⬇️ Export Report")
        col_ex1, col_ex2 = st.columns(2)

        with col_ex1:
            if EXCEL_AVAILABLE:
                excel_data = export_to_excel(df_rep, r_dept, date_str)
                st.download_button(
                    label="📥 Download Excel (.xlsx)",
                    data=excel_data,
                    file_name=f"{r_dept}_{date_str}_attendance.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                st.warning("Install openpyxl for Excel export: `pip install openpyxl`")

        with col_ex2:
            if PDF_AVAILABLE:
                pdf_data = export_to_pdf(df_rep, r_dept, date_str)
                st.download_button(
                    label="📥 Download PDF",
                    data=pdf_data,
                    file_name=f"{r_dept}_{date_str}_attendance.pdf",
                    mime="application/pdf"
                )
            else:
                st.warning("Install reportlab for PDF export: `pip install reportlab`")

    else:
        st.info(f"No attendance recorded for **{r_dept}** on **{date_str}**.")


@st.cache_resource
def load_resources():
    recognizers    = {}
    student_lookup = {}
    trained_set    = set()

    for dept in DEPARTMENTS:
        for section in SECTIONS:
            model_path = BASE_DIR / dept / section / "trainer.yml"
            if model_path.exists():
                rec = create_lbph_recognizer()
                if rec:
                    rec.read(str(model_path))
                    recognizers[f"{dept}_{section}"] = rec

            csv_path = BASE_DIR / dept / section / "students.csv"
            if csv_path.exists():
                try:
                    df = pd.read_csv(csv_path, dtype=str)
                    for _, r in df.iterrows():
                        roll = str(r.get("roll_number", "")).strip()
                        if roll:
                            student_lookup[(f"{dept}_{section}", int(roll))] = {
                                "roll":    roll,
                                "name":    r.get("name", ""),
                                "dept":    dept,
                                "section": section,
                                "email":   r.get("email", "")
                            }
                            try:
                                pass  # Labels are scoped by department and section.
                            except Exception:
                                pass
                except Exception:
                    pass

            faces_dir = BASE_DIR / dept / section / "faces"
            if faces_dir.exists():
                for s_dir in faces_dir.iterdir():
                    if s_dir.is_dir():
                        trained_set.add(s_dir.name)
                        try:
                            trained_set.add(int(s_dir.name))
                        except Exception:
                            pass

    return recognizers, student_lookup, trained_set


def run_attendance_loop(video_placeholder, log_placeholder):
    recognizers, student_lookup, trained_set = load_resources()

    if not recognizers:
        st.error("❌ No trained models found. Please go to Training page and train first.")
        st.session_state.attendance_active = False
        return

    cap          = open_camera()
    if cap is None:
        st.session_state.attendance_active = False
        return
    face_cascade = get_face_detector()

    attendance_session = set()

    try:
        while st.session_state.attendance_active:
            ret, frame = cap.read()
            if not ret:
                st.session_state.attendance_active = False
                st.error("Camera stopped delivering frames. Close other camera apps and restart recognition.")
                break
    
            gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            faces = detect_faces(gray, face_cascade)
    
            if not faces:
                log_placeholder.empty()

            pname, pstart, pend = get_current_period()
    
            for (x, y, w, h) in faces:
                face_roi = gray[y:y+h, x:x+w]
                try:
                    face_roi = prepare_face(face_roi)
                except Exception:
                    continue
    
                best_confidence = 1000
                best_match      = None
    
                for key, recognizer in recognizers.items():
                    try:
                        label, confidence = recognizer.predict(face_roi)
                        if confidence < 55 and confidence < best_confidence:
                            if label in trained_set or str(label) in trained_set:
                                best_confidence = confidence
                                best_match      = (key, label)
                    except Exception:
                        pass
    
                color     = (0, 0, 255)
                name_text = "Unknown"
    
                if best_match is not None:
                    info = student_lookup.get(best_match)
                    if info:
                        name_text = f"{info['name']} ({info['roll']})"
                        color     = (0, 255, 0)
    
                        if pname and "Break" not in pname and not st.session_state.get("preview_only", False):
                            status = mark_attendance(info, pname, pstart)
                            if info["roll"] not in attendance_session:
                                if status == "Late":
                                    log_placeholder.warning(f"⏰ Late: {name_text} for {pname}")
                                else:
                                    log_placeholder.success(f"✅ Present: {name_text} for {pname}")
                                attendance_session.add(info["roll"])
    
                cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                cv2.putText(frame, name_text, (x, y-10), cv2.FONT_HERSHEY_SIMPLEX, 0.8, color, 2)
    
            frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            video_placeholder.image(frame_rgb, channels="RGB")
    finally:
        cap.release()


def mark_attendance(info, period, period_start):
    date_str    = datetime.now().strftime("%Y-%m-%d")
    reports_dir = BASE_DIR / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    csv_path = reports_dir / f"{info['dept']}_{date_str}_attendance.csv"

    if csv_path.exists():
        df = pd.read_csv(csv_path, dtype={"roll_number": str})
    else:
        df = pd.DataFrame(columns=["date", "roll_number", "name", "section", "period", "sign_in_time", "status"])

    mask = (
        (df["date"] == date_str) &
        (df["roll_number"].astype(str) == str(info["roll"])) &
        (df["period"] == period)
    )

    if not mask.any():
        now        = datetime.now()
        sign_in_str = now.strftime("%Y-%m-%d %H:%M:%S")
        status     = get_attendance_status(sign_in_str, period_start)

        new_row = {
            "date":        date_str,
            "roll_number": info["roll"],
            "name":        info["name"],
            "section":     info["section"],
            "period":      period,
            "sign_in_time": sign_in_str,
            "status":      status
        }
        df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
        df.to_csv(csv_path, index=False)
        return status

    # Return existing status if already marked
    existing = df[mask]["status"].values
    return existing[0] if len(existing) > 0 else "Present"


# -----------------------
# Excel Export
# -----------------------
def export_to_excel(df, dept, date_str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Title
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{dept} Department — Attendance Report — {date_str}"
    ws["A1"].font      = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row
    headers = ["Date", "Roll Number", "Name", "Section", "Period", "Sign-In Time", "Status"]
    header_fill = PatternFill(start_color="2F75B6", end_color="2F75B6", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border    = border

    # Data rows
    present_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    late_fill    = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

    for row_idx, row in df.iterrows():
        status = str(row.get("status", "Present"))
        fill   = present_fill if status == "Present" else late_fill

        for col_idx, header in enumerate(headers, start=1):
            col_key = header.lower().replace(" ", "_").replace("-", "_")
            # map header to df column
            col_map = {
                "date": "date", "roll_number": "roll_number", "name": "name",
                "section": "section", "period": "period",
                "sign_in_time": "sign_in_time", "status": "status"
            }
            val  = row.get(col_map.get(col_key, col_key), "")
            cell = ws.cell(row=row_idx + 3, column=col_idx, value=val)
            cell.fill      = fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center")

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# -----------------------
# PDF Export
# -----------------------
def export_to_pdf(df, dept, date_str):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1, fontSize=16)
    elements.append(Paragraph(f"{dept} Department — Attendance Report", title_style))
    elements.append(Paragraph(f"Date: {date_str}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    # Summary
    present_count = len(df[df["status"] == "Present"]) if "status" in df.columns else 0
    late_count    = len(df[df["status"] == "Late"])    if "status" in df.columns else 0
    elements.append(Paragraph(f"Total: {len(df)} | Present: {present_count} | Late: {late_count}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    # Table
    headers = ["Roll No", "Name", "Section", "Period", "Sign-In Time", "Status"]
    col_map = ["roll_number", "name", "section", "period", "sign_in_time", "status"]
    data    = [headers]

    for _, row in df.iterrows():
        data.append([str(row.get(c, "")) for c in col_map])

    col_widths = [70, 100, 60, 60, 120, 60]
    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ("BACKGROUND",  (0, 0), (-1, 0),  colors.HexColor("#2F75B6")),
        ("TEXTCOLOR",   (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",    (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0, 0), (-1, 0),  10),
        ("ALIGN",       (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME",    (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",    (0, 1), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8F9FA"), colors.white]),
        ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("TOPPADDING",  (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]

    # Highlight Late rows in yellow
    for i, row in enumerate(df.itertuples(), start=1):
        if hasattr(row, "status") and row.status == "Late":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF3CD")))

    t.setStyle(TableStyle(style_cmds))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)
    return buf.read()


# -----------------------
# Page: Email Alerts
# -----------------------
def page_email_alerts():
    st.header("📧 Email Alerts Configuration")

    st.info("""
    Configure your Gmail SMTP settings here (smtp.gmail.com, SSL port 465).
    Once configured, you can send absence alerts to students/parents.
    """)

    with st.form("email_config_form"):
        st.subheader("SMTP Settings")
        sender_email = st.text_input("Sender Email (Gmail)", placeholder="yourschool@gmail.com")
        app_password = st.text_input(
            "App Password",
            type="password",
            help="Use Gmail App Password (not your regular password). Enable 2FA first, then generate app password."
        )
        st.form_submit_button("💾 Save Settings (stored in session only)")

    if sender_email:
        st.session_state["sender_email"] = sender_email
    if app_password:
        st.session_state["app_password"] = app_password.replace(" ", "")

    st.markdown("---")
    st.subheader("📤 Send Absence Alert")

    col1, col2 = st.columns(2)
    with col1:
        alert_dept = st.selectbox("Department", DEPARTMENTS, key="alert_dept")
        alert_date = st.date_input("Date", datetime.now()).strftime("%Y-%m-%d")
    with col2:
        alert_period = st.selectbox("Period", [f"Period{i}" for i in range(1, 13)], key="alert_period")

    if st.button("📬 Send Absence Alerts"):
        send_absence_alerts(alert_dept, alert_date, alert_period)

    st.markdown("---")
    st.subheader("📤 Send Daily Summary to Teacher")
    teacher_email = st.text_input("Teacher Email")
    summary_dept  = st.selectbox("Department for Summary", DEPARTMENTS, key="sum_dept")
    summary_date  = st.date_input("Summary Date", datetime.now(), key="sum_date").strftime("%Y-%m-%d")

    if st.button("📩 Send Summary Email"):
        send_summary_email(teacher_email, summary_dept, summary_date)


def send_absence_alerts(dept, date_str, period):
    sender_email = st.session_state.get("sender_email", "")
    app_password = st.session_state.get("app_password", "")

    if not sender_email or not app_password:
        st.error("Please configure email settings first.")
        return

    # Load all students
    present_rolls = set()
    report_path   = BASE_DIR / "reports" / f"{dept}_{date_str}_attendance.csv"

    if report_path.exists():
        df_rep = pd.read_csv(report_path, dtype={"roll_number": str})
        present_in_period = df_rep[df_rep["period"] == period]
        present_rolls     = set(present_in_period["roll_number"].astype(str).tolist())

    # Get all students in dept
    all_students = []
    for section in SECTIONS:
        csv_path = BASE_DIR / dept / section / "students.csv"
        if csv_path.exists():
            df_s = pd.read_csv(csv_path, dtype=str)
            for _, r in df_s.iterrows():
                all_students.append({
                    "roll": str(r.get("roll_number", "")),
                    "name": r.get("name", ""),
                    "email": r.get("email", "")
                })

    absent_students = [s for s in all_students if s["roll"] not in present_rolls and s["email"]]

    if not absent_students:
        st.info("No absent students found (or no emails configured).")
        return

    sent = 0
    failed = 0
    for student in absent_students:
        try:
            msg            = MIMEMultipart()
            msg["From"]    = sender_email
            msg["To"]      = student["email"]
            msg["Subject"] = f"Attendance Alert — {date_str}"

            body = f"""
Dear Parent/Guardian,

This is to inform you that your ward {student['name']} (Roll: {student['roll']})
was ABSENT during {period} on {date_str} for {dept} department.

Please ensure regular attendance.

Regards,
College Attendance System
            """.strip()

            msg.attach(MIMEText(body, "plain"))
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20) as server:
                server.login(sender_email, app_password)
                server.sendmail(sender_email, student["email"], msg.as_string())
            sent += 1
        except Exception as e:
            failed += 1

    st.success(f"✅ Sent {sent} absence alerts. Failed: {failed}")


def send_summary_email(teacher_email, dept, date_str):
    sender_email = st.session_state.get("sender_email", "")
    app_password = st.session_state.get("app_password", "")

    if not sender_email or not app_password:
        st.error("Please configure email settings first.")
        return

    if not teacher_email:
        st.error("Please enter teacher email.")
        return

    report_path = BASE_DIR / "reports" / f"{dept}_{date_str}_attendance.csv"
    if not report_path.exists():
        st.error("No attendance report found for selected date/department.")
        return

    df_rep = pd.read_csv(report_path, dtype={"roll_number": str})

    try:
        msg            = MIMEMultipart()
        msg["From"]    = sender_email
        msg["To"]      = teacher_email
        msg["Subject"] = f"Daily Attendance Summary — {dept} — {date_str}"

        present_count = len(df_rep[df_rep["status"] == "Present"]) if "status" in df_rep.columns else 0
        late_count    = len(df_rep[df_rep["status"] == "Late"])    if "status" in df_rep.columns else 0
        total         = len(df_rep)

        body = f"""
Dear Teacher,

Here is the attendance summary for {dept} Department on {date_str}:

Total Records : {total}
Present       : {present_count}
Late          : {late_count}
Absent        : (Not yet auto-tracked in this version)

Please find the detailed CSV report attached.

Regards,
College Attendance System
        """.strip()

        msg.attach(MIMEText(body, "plain"))

        # Attach CSV
        with open(report_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f"attachment; filename={report_path.name}")
            msg.attach(part)

        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20) as server:
            server.login(sender_email, app_password)
            server.sendmail(sender_email, teacher_email, msg.as_string())

        st.success(f"✅ Summary email sent to {teacher_email}!")

    except Exception as e:
        st.error(f"Failed to send email: {e}")


# -----------------------
# Main App Structure
# -----------------------
def main():
    st.set_page_config(page_title="Face Attendance System", page_icon="🎓", layout="wide")

    # Check for missing optional libraries
    missing = []
    if not EXCEL_AVAILABLE:
        missing.append("`pip install openpyxl`")
    if not PDF_AVAILABLE:
        missing.append("`pip install reportlab`")
    if missing:
        st.sidebar.warning("⚠️ Optional libraries missing:\n" + "\n".join(missing))

    st.sidebar.title("🎓 Face Attendance")
    st.sidebar.number_input("Camera number", min_value=0, max_value=5, value=0, step=1, key="camera_index")
    st.sidebar.caption("0 is usually the built-in webcam. Stop capture before switching cameras.")
    st.sidebar.markdown("---")
    page = st.sidebar.radio("Navigate", ["📋 Attendance", "🎓 Training", "📧 Email Alerts"])

    if page != "🎓 Training":
        st.session_state.training_active = False
    if page != "📋 Attendance":
        st.session_state.attendance_active = False

    if page == "🎓 Training":
        page_training()
    elif page == "📧 Email Alerts":
        page_email_alerts()
    else:
        page_attendance()

    st.sidebar.markdown("---")
    st.sidebar.caption("Built with Streamlit + OpenCV")


if __name__ == "__main__":
    main()
