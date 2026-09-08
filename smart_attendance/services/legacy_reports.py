import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import inch
def export_to_excel(df, dept, date_str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Title
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{dept} Department — Attendance Report — {date_str}"
    ws["A1"].font      = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row
    headers = ["Date", "Roll Number", "Name", "Section", "Period", "Sign-In Time", "Status"]
    header_fill = PatternFill(start_color="2F75B6", end_color="2F75B6", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border    = border

    # Data rows
    present_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    late_fill    = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

    for row_idx, row in df.iterrows():
        status = str(row.get("status", "Present"))
        fill   = present_fill if status == "Present" else late_fill

        for col_idx, header in enumerate(headers, start=1):
            col_key = header.lower().replace(" ", "_").replace("-", "_")
            # map header to df column
            col_map = {
                "date": "date", "roll_number": "roll_number", "name": "name",
                "section": "section", "period": "period",
                "sign_in_time": "sign_in_time", "status": "status"
            }
            val  = row.get(col_map.get(col_key, col_key), "")
            cell = ws.cell(row=row_idx + 3, column=col_idx, value=val)
            cell.fill      = fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center")

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

def export_to_pdf(df, dept, date_str):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1, fontSize=16)
    elements.append(Paragraph(f"{dept} Department — Attendance Report", title_style))
    elements.append(Paragraph(f"Date: {date_str}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    # Summary
    present_count = len(df[df["status"] == "Present"]) if "status" in df.columns else 0
    late_count    = len(df[df["status"] == "Late"])    if "status" in df.columns else 0
    elements.append(Paragraph(f"Total: {len(df)} | Present: {present_count} | Late: {late_count}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    # Table
    headers = ["Roll No", "Name", "Section", "Period", "Sign-In Time", "Status"]
    col_map = ["roll_number", "name", "section", "period", "sign_in_time", "status"]
    data    = [headers]

    for _, row in df.iterrows():
        data.append([str(row.get(c, "")) for c in col_map])

    col_widths = [70, 100, 60, 60, 120, 60]
    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ("BACKGROUND",  (0, 0), (-1, 0),  colors.HexColor("#2F75B6")),
        ("TEXTCOLOR",   (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",    (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0, 0), (-1, 0),  10),
        ("ALIGN",       (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME",    (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",    (0, 1), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8F9FA"), colors.white]),
        ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("TOPPADDING",  (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]

    # Highlight Late rows in yellow
    for i, row in enumerate(df.itertuples(), start=1):
        if hasattr(row, "status") and row.status == "Late":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF3CD")))

    t.setStyle(TableStyle(style_cmds))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)
    return buf.read()