import streamlit as st
import cv2
import os
import sys
import numpy as np
import pandas as pd
from pathlib import Path
from PIL import Image
from datetime import datetime, timedelta
import time
import traceback
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
import io
import shutil
import threading
from face_pipeline import prepare_face, valid_sample, detect_faces
from smart_attendance import ui_theme as ui
from smart_attendance.database.migrations import import_legacy
from smart_attendance.database.db import connect
from smart_attendance.services.attendance_service import (
    finalize_due, record_observation, refresh_report, rules as attendance_rules, start_class, sync_csv,
)
from smart_attendance.services.liveness_service import (
    BlinkChallengeProvider, FaceTracker, LivenessStatus,
)
from smart_attendance.services.timetable_runtime import current_classes, fallback_context, student_semesters
from smart_attendance.services.audit_service import record_system

# Try importing optional libraries
try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.units import inch
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Check if running with streamlit run
try:
    from streamlit.runtime.scriptrunner import get_script_run_ctx
    if not get_script_run_ctx():
        print("\n\n" + "="*60)
        print("ERROR: You are running this script incorrectly!")
        print("Please do NOT run 'python streamlit_app.py'.")
        print("Instead, run: streamlit run streamlit_app.py")
        print("="*60 + "\n\n")
except ImportError:
    pass

# -----------------------
# Configuration & Helpers
# -----------------------
BASE_DIR = Path(__file__).resolve().parent / "dataset" / "college_faces"
BASE_DIR.mkdir(parents=True, exist_ok=True)
DEPARTMENTS = ["AIML", "CSE", "Civil", "ECE", "MECH"]
SECTIONS = ["A", "B", "C"]

def create_lbph_recognizer():
    if hasattr(cv2, "face") and hasattr(cv2.face, "LBPHFaceRecognizer_create"):
        return cv2.face.LBPHFaceRecognizer_create()
    try:
        return cv2.face.LBPHFaceRecognizer_create()
    except Exception:
        st.error("LBPH recognizer not available. Please install opencv-contrib-python.")
        return None

def get_schedule():
    def t(h, m): return datetime.strptime(f"{h}:{m}", "%H:%M").time()
    return [
        ("Period1",  t(8,30),  t(9,30)),
        ("Period2",  t(9,30),  t(10,30)),
        ("Break1",   t(10,30), t(11,0)),
        ("Period3",  t(11,0),  t(12,0)),
        ("Period4",  t(12,0),  t(13,0)),
        ("Break2",   t(13,0),  t(13,40)),
        ("Period5",  t(13,40), t(14,40)),
        ("Period6",  t(14,40), t(15,40)),
        ("Period7",  t(18,0),  t(19,0)),
        ("Period8",  t(19,0),  t(20,0)),
        ("Period9",  t(20,0),  t(21,0)),
        ("Period10", t(21,0),  t(22,0)),
        ("Period11", t(22,0),  t(23,0)),
        ("Period12", t(23,0),  t(23,59)),
    ]

def get_current_period():
    now = datetime.now().time()
    for name, start, end in get_schedule():
        if start <= now < end:
            return name, start, end
    return None, None, None

def get_attendance_status(sign_in_time_str, period_start):
    """
    Determines if student is Present or Late.
    Late = arrived more than 10 minutes after period started.
    """
    try:
        sign_in_time = datetime.strptime(sign_in_time_str, "%Y-%m-%d %H:%M:%S").time()
        period_start_dt = datetime.combine(datetime.today(), period_start)
        sign_in_dt = datetime.combine(datetime.today(), sign_in_time)
        diff_minutes = (sign_in_dt - period_start_dt).total_seconds() / 60
        if diff_minutes > 10:
            return "Late"
        return "Present"
    except Exception:
        return "Present"

# -----------------------
# Page: Training
# -----------------------
def open_camera():
    from smart_attendance.services.camera_service import manager, CameraError
    import uuid
    owner = uuid.uuid4().hex
    try:
        manager.start(owner, int(st.session_state.get("camera_index", 0)))
    except CameraError as exc:
        st.error(str(exc))
        return None
    class Lease:
        def read(self):
            try:
                for _ in range(30):
                    frame = manager.read(owner)
                    if frame is not None:
                        return True, frame
                    time.sleep(.05)
            except CameraError:
                return False, None
            return False, None
        def release(self):
            manager.close(owner)
    return Lease()


def browser_camera_enabled():
    return st.session_state.get("camera_source", "Browser camera") == "Browser camera"


def browser_camera_component(key, callback):
    """Render a browser-owned camera stream for hosted deployments."""
    try:
        from streamlit_webrtc import WebRtcMode, webrtc_streamer
    except ImportError:
        st.error("Browser camera support is not installed. Install the current requirements and restart the app.")
        return None
    return webrtc_streamer(
        key=key,
        mode=WebRtcMode.SENDRECV,
        video_frame_callback=callback,
        media_stream_constraints={"video": True, "audio": False},
        rtc_configuration={"iceServers": [{"urls": ["stun:stun.l.google.com:19302"]}]},
        async_processing=True,
        media_toggle_controls=False,
    )


def browser_camera_test():
    detector = get_face_detector()

    def annotate(video_frame):
        import av
        frame = video_frame.to_ndarray(format="bgr24")
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        faces = list(detector.detectMultiScale(gray, 1.1, 8, minSize=(80, 80)))
        for x, y, w, h in faces:
            cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 180, 70), 2)
        label = "CAMERA READY" if faces else "CAMERA READY | NO FACE DETECTED"
        cv2.putText(frame, label, (18, 30), cv2.FONT_HERSHEY_SIMPLEX, .65, (0, 180, 70), 2)
        return av.VideoFrame.from_ndarray(frame, format="bgr24")

    context = browser_camera_component("smartattend-browser-test", annotate)
    if context is not None:
        if context.state.playing:
            st.success("Browser camera connected successfully.")
        else:
            st.info("Click START and choose Allow when the browser requests camera permission.")


def browser_face_enrollment(faces_dir, capture_dir, roll, name, dept, section):
    detector = get_face_detector()
    state = {"count": 0, "last_face": None, "last_saved": 0.0, "complete": False}
    lock = threading.Lock()
    required_samples = 50

    def capture(video_frame):
        import av
        frame = video_frame.to_ndarray(format="bgr24")
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        faces = detect_faces(gray, detector)
        with lock:
            if len(faces) == 1 and not state["complete"]:
                x, y, w, h = faces[0]
                face_img = gray[y:y+h, x:x+w]
                try:
                    normalized = prepare_face(cv2.resize(face_img, (200, 200)))
                    now_saved = time.monotonic()
                    changed = state["last_face"] is None or float(
                        np.mean(cv2.absdiff(normalized, state["last_face"]))
                    ) >= 1.8
                    if changed and now_saved - state["last_saved"] >= .16:
                        target = capture_dir / f"face_{state['count']}.jpg"
                        if cv2.imwrite(str(target), face_img):
                            state["count"] += 1
                            state["last_face"] = normalized
                            state["last_saved"] = now_saved
                except (ValueError, cv2.error):
                    pass

                cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 180, 70), 2)

                if state["count"] >= required_samples:
                    if faces_dir.exists():
                        backup = (
                            BASE_DIR.parent / "face_backups" / datetime.now().strftime("%Y%m%d_%H%M%S")
                            / dept / section / roll
                        )
                        backup.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copytree(faces_dir, backup)
                        shutil.rmtree(faces_dir)
                    capture_dir.replace(faces_dir)
                    from train_faces import train_all
                    training_report = train_all(BASE_DIR)
                    profile = next(
                        (row for row in training_report
                         if str(row.get("roll")) == str(roll)
                         and str(row.get("section", "")).replace("\\", "/") == f"{dept}/{section}"),
                        None,
                    )
                    training_status = str((profile or {}).get("status", "Training did not produce a profile"))
                    (faces_dir / ".enrollment_complete").write_text(training_status, encoding="utf-8")
                    record_system(
                        "FACE ENROLLED", f"{dept}:{section}:{roll}",
                        {
                            "department": dept, "section": section, "roll": roll,
                            "sample_count": state["count"], "training_status": training_status,
                        },
                    )
                    state["complete"] = True

            if len(faces) > 1:
                message, color = "ONE FACE ONLY", (0, 120, 230)
            elif state["complete"]:
                message, color = "CAPTURE COMPLETE - CLICK STOP", (0, 180, 70)
            else:
                message, color = f"CAPTURED {state['count']} / {required_samples}", (0, 180, 70)
            cv2.putText(frame, message, (18, 30), cv2.FONT_HERSHEY_SIMPLEX, .65, color, 2)
        return av.VideoFrame.from_ndarray(frame, format="bgr24")

    context = browser_camera_component(f"smartattend-enroll-{dept}-{section}-{roll}", capture)
    if context is not None:
        if context.state.playing:
            st.info("Capture is running. Slowly vary your angle and keep both eyes visible.")
        else:
            st.info("Click START and allow browser camera access. Stop when the video says capture complete.")


def get_face_detector():
    detector = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
    if detector.empty():
        raise RuntimeError("Face detector missing. Install the packages in requirements.txt and restart the app.")
    return detector


def page_training():
    ui.heading("STUDENT DIRECTORY", "Enroll. Capture. Recognize.", "Register students and build reliable face profiles for your classroom.")
    st.info("Capture one student at a time in even light. Keep eyes visible, face the camera, and vary your angle slightly. Blurry or poorly detected samples are skipped.")

    col1, col2, col3 = st.columns(3)
    with col1:
        roll = st.text_input("Roll Number (numbers only)")
        name = st.text_input("Name")
    with col2:
        dept = st.selectbox("Department", DEPARTMENTS)
        section = st.selectbox("Section", SECTIONS)
    with col3:
        semester = st.number_input("Semester", min_value=1, max_value=12, value=5, step=1)

    if "training_active" not in st.session_state:
        st.session_state.training_active = False

    action_start, action_stop, action_test = st.columns(3)
    with action_start:
        start_btn = st.button("Start capture", disabled=st.session_state.training_active, type="primary", width="stretch")
    with action_stop:
        stop_btn = st.button("Stop capture", disabled=not st.session_state.training_active, width="stretch")
    with action_test:
        test_btn = st.button("Test camera", disabled=st.session_state.training_active, width="stretch")

    if start_btn:
        if not roll or not name:
            st.error("Please enter Roll Number and Name.")
        else:
            try:
                int(roll)
            except ValueError:
                st.error("Roll Number must be a plain integer (e.g. 101, 12345).")
                return
            st.session_state.training_active   = True
            st.session_state.training_roll     = roll
            st.session_state.training_name     = name
            st.session_state.training_dept     = dept
            st.session_state.training_section  = section
            st.session_state.training_semester = int(semester)
            st.rerun()

    if stop_btn:
        st.session_state.training_active = False
        st.rerun()

    video_placeholder = st.empty()
    progress_bar      = st.progress(0)
    status_text       = st.empty()

    if test_btn:
        st.session_state.browser_camera_test_active = browser_camera_enabled()
        if not browser_camera_enabled():
            cap = open_camera()
            if cap is not None:
                try:
                    ok, frame = cap.read()
                    if not ok or frame is None:
                        st.error("The camera opened but did not return a frame. Try another camera number.")
                    else:
                        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
                        faces = detect_faces(gray, get_face_detector())
                        frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                        video_placeholder.image(frame_rgb, channels="RGB")
                        if len(faces) == 1:
                            st.success("Camera is ready. One enrollment-quality face was detected.")
                        elif len(faces) > 1:
                            st.warning("Camera is working, but multiple faces are visible. Keep only one student in frame for capture.")
                        else:
                            st.info("Camera is working. No enrollment-quality face is currently centered with both eyes visible.")
                finally:
                    cap.release()

    if st.session_state.get("browser_camera_test_active") and not st.session_state.training_active:
        browser_camera_test()

    if st.session_state.training_active:
        roll    = st.session_state.training_roll
        dept    = st.session_state.training_dept
        section = st.session_state.training_section
        semester = st.session_state.training_semester
        name    = st.session_state.training_name

        faces_root = BASE_DIR / dept / section / "faces"
        faces_dir = faces_root / roll
        capture_dir = faces_root / f".capture_{roll}"
        faces_root.mkdir(parents=True, exist_ok=True)
        completion_marker = faces_dir / ".enrollment_complete"
        if browser_camera_enabled() and completion_marker.exists():
            training_status = completion_marker.read_text(encoding="utf-8")
            completion_marker.unlink()
            st.session_state.training_active = False
            load_resources.clear()
            if training_status.startswith("Trained"):
                st.success(f"Capture and model training completed for {name} ({roll}).")
            else:
                st.error(f"Capture completed, but validation failed: {training_status}. Please re-enroll with clearer angles and even lighting.")
            return
        if capture_dir.exists():
            shutil.rmtree(capture_dir)
        capture_dir.mkdir(parents=True)

        # Update students.csv
        dept_dir = BASE_DIR / dept / section
        dept_dir.mkdir(parents=True, exist_ok=True)
        csv_path = dept_dir / "students.csv"

        if csv_path.exists():
            df = pd.read_csv(csv_path, dtype=str)
        else:
            df = pd.DataFrame(columns=["roll_number", "name", "branch", "section", "semester", "email"])

        if "semester" not in df.columns:
            df["semester"] = "1"

        is_new_student = not df["roll_number"].astype(str).eq(str(roll)).any()
        if is_new_student:
            new_row = {
                "roll_number": roll, "name": name, "branch": dept,
                "section": section, "semester": str(semester), "email": "",
            }
            df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
            df.to_csv(csv_path, index=False)
            record_system(
                "STUDENT REGISTERED", f"{dept}:{section}:{roll}",
                {"department": dept, "section": section, "semester": semester, "roll": roll},
            )

        if browser_camera_enabled():
            st.session_state.browser_camera_test_active = False
            browser_face_enrollment(faces_dir, capture_dir, roll, name, dept, section)
            return

        cap          = open_camera()
        if cap is None:
            st.session_state.training_active = False
            return
        face_cascade = get_face_detector()

        sample_count    = 0
        required_samples = 50
        last_saved_face = None
        last_saved_at = 0.0

        try:
            while st.session_state.training_active and sample_count < required_samples:
                ret, frame = cap.read()
                if not ret:
                    st.session_state.training_active = False
                    st.error("Camera stopped delivering frames. Restart capture.")
                    break
    
                gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
                faces = detect_faces(gray, face_cascade)
    
                if len(faces) > 1:
                    status_text.warning("More than one face detected. Keep only the student being enrolled in frame.")
                for (x, y, w, h) in (faces if len(faces) == 1 else []):
                    cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 255, 0), 2)
                    face_img = gray[y:y+h, x:x+w]
                    try:
                        face_img = cv2.resize(face_img, (200, 200))
                        normalized = prepare_face(face_img)
                        now_saved = time.monotonic()
                        different_enough = (
                            last_saved_face is None
                            or float(np.mean(cv2.absdiff(normalized, last_saved_face))) >= 1.8
                        )
                        if now_saved - last_saved_at < 0.16 or not different_enough:
                            continue
                        if cv2.imwrite(str(capture_dir / f"face_{sample_count}.jpg"), face_img):
                            sample_count += 1
                            last_saved_face = normalized
                            last_saved_at = now_saved
                    except Exception:
                        continue
    
                progress_bar.progress(min(sample_count / required_samples, 1.0))
                status_text.text(
                    f"Captured: {sample_count} / {required_samples} · Slowly vary your angle while keeping both eyes visible."
                )
    
                frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                video_placeholder.image(frame_rgb, channels="RGB")
    
                if sample_count >= required_samples:
                    st.session_state.training_active = False
                    cap.release()
                    if faces_dir.exists():
                        backup = (
                            BASE_DIR.parent / "face_backups" / datetime.now().strftime("%Y%m%d_%H%M%S")
                            / dept / section / roll
                        )
                        backup.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copytree(faces_dir, backup)
                        shutil.rmtree(faces_dir)
                    capture_dir.replace(faces_dir)
                    record_system(
                        "FACE ENROLLED", f"{dept}:{section}:{roll}",
                        {"department": dept, "section": section, "roll": roll, "sample_count": sample_count},
                    )
                    st.success(f"✅ Capture completed for {name} ({roll})! Now click 'Train Models'.")
                    st.rerun()
        finally:
            cap.release()

    st.markdown("---")
    if st.button("Train all models"):
        train_models()

    # ---- Email management ----
    st.markdown("---")
    st.subheader("Student & parent contacts")
    st.info("Add parent/student email addresses here to enable absence alerts.")

    em_dept    = st.selectbox("Department", DEPARTMENTS, key="em_dept")
    em_section = st.selectbox("Section",    SECTIONS,    key="em_section")

    csv_path = BASE_DIR / em_dept / em_section / "students.csv"
    if csv_path.exists():
        df_em = pd.read_csv(csv_path, dtype=str)
        if "email" not in df_em.columns:
            df_em["email"] = ""

        edited_df = st.data_editor(
            df_em[["roll_number", "name", "email"]],
            width="stretch",
            key="email_editor"
        )

        if st.button("Save email addresses"):
            df_em["email"] = edited_df["email"].values
            df_em.to_csv(csv_path, index=False)
            st.success("Emails saved!")
    else:
        st.warning("No students registered for this department/section yet.")


def train_models():
    from train_faces import train_all
    with st.spinner("Checking samples and training models..."):
        report = train_all(BASE_DIR)
    load_resources.clear()
    record_system(
        "FACE MODELS TRAINED", "all-sections",
        {"profiles": len(report), "validated": sum(str(row.get("status", "")).startswith("Trained") for row in report)},
    )
    st.dataframe(pd.DataFrame(report), width="stretch")
    st.success("Training finished. See accepted samples and validation results above.")


# -----------------------
# Page: Attendance
# -----------------------
def page_attendance():
    ui.heading("LIVE ATTENDANCE", "Every presence counts.", "Recognize students, record attendance, and keep your classroom in view.")
    ui.overview(BASE_DIR)
    st.checkbox("Preview only (do not record attendance)", key="preview_only")

    scheduled = current_classes()
    if scheduled:
        selected_context = st.selectbox(
            "Current scheduled class", scheduled, format_func=lambda item: item.label
        )
        st.caption(f"Faculty: {selected_context.faculty or 'Not assigned'} · Period: {selected_context.period}")
    else:
        st.caption("No configured class is active. Select a class scope to use the existing period schedule.")
        scope_left, scope_right = st.columns(2)
        with scope_left:
            fallback_department = st.selectbox("Fallback department", DEPARTMENTS, key="live_department")
        with scope_right:
            fallback_section = st.selectbox("Fallback section", SECTIONS, key="live_section")
        semesters = student_semesters(fallback_department, fallback_section)
        if semesters:
            fallback_semester = st.selectbox(
                "Fallback semester", semesters, index=len(semesters) - 1, key="live_semester",
                help="Attendance is created only for active students in this semester.",
            )
            selected_context = fallback_context(
                fallback_department, fallback_section, get_schedule(), semester=fallback_semester,
            )
        else:
            st.warning(f"No active students are registered in {fallback_department}-{fallback_section}.")
            selected_context = None

    if "attendance_active" not in st.session_state:
        st.session_state.attendance_active = False

    if selected_context is not None:
        expected_model = BASE_DIR / selected_context.department / selected_context.section / "trainer.yml"
        recognition_ready = expected_model.exists()
        if not recognition_ready:
            st.warning(
                f"Attendance is waiting for a trained face model for "
                f"{selected_context.department}-{selected_context.section}. Open Students & training, "
                "complete face capture, and wait for model validation to finish."
            )
    else:
        recognition_ready = any(BASE_DIR.glob("*/*/trainer.yml"))

    col1, col2, col3 = st.columns(3)
    with col1:
        start_btn = st.button(
            "Start recognition",
            disabled=st.session_state.attendance_active or not recognition_ready,
            type="primary",
            width="stretch",
        )
    with col2:
        stop_btn  = st.button("Stop recognition",  disabled=not st.session_state.attendance_active, width="stretch")
    with col3:
        if st.button("Reload data", width="stretch"):
            st.cache_resource.clear()
            st.success("Reloaded!")

    if selected_context:
        st.info(
            f"Current class: **{selected_context.subject}** · "
            f"{selected_context.department}-{selected_context.section} · "
            f"{selected_context.starts_at:%H:%M}–{selected_context.ends_at:%H:%M}"
        )
    else:
        st.warning("No active class or fallback period. Preview remains available, but attendance cannot be recorded.")

    video_placeholder = st.empty()
    log_placeholder   = st.empty()

    if stop_btn:
        st.session_state.attendance_active = False
        st.session_state.pop("attendance_context", None)
        st.rerun()

    if start_btn:
        if not st.session_state.preview_only and selected_context is None:
            st.error("Select or configure an active class before recording attendance.")
        else:
            if selected_context is not None and not st.session_state.preview_only:
                try:
                    st.session_state.attendance_session_id = start_class(selected_context)
                    sync_csv(selected_context.starts_at.date(), selected_context.department, BASE_DIR)
                except ValueError as exc:
                    st.error(str(exc))
                    return
            st.session_state.attendance_context = selected_context
            st.session_state.attendance_active = True
            st.rerun()

    if st.session_state.attendance_active:
        run_attendance_loop(
            video_placeholder,
            log_placeholder,
            st.session_state.get("attendance_context"),
        )

    if not st.session_state.attendance_active:
        ui.camera_idle()

    # ---- Reports Section ----
    st.markdown("---")
    st.subheader("Attendance reports")

    r_dept = st.selectbox("Select Department", DEPARTMENTS, key="rep_dept")
    date_str = st.date_input("Select Date", datetime.now()).strftime("%Y-%m-%d")

    report_path = refresh_report(datetime.strptime(date_str, "%Y-%m-%d").date(), r_dept, BASE_DIR)

    if report_path and report_path.exists():
        df_rep = pd.read_csv(report_path, dtype={"roll_number": str})
        filter_columns = st.columns(3)
        with filter_columns[0]:
            subject_options = ["All"] + sorted(df_rep["subject"].dropna().astype(str).unique().tolist()) if "subject" in df_rep else ["All"]
            report_subject = st.selectbox("Subject", subject_options, key="report_subject")
        with filter_columns[1]:
            section_options = ["All"] + sorted(df_rep["section"].dropna().astype(str).unique().tolist())
            report_section = st.selectbox("Report section", section_options, key="report_section")
        with filter_columns[2]:
            status_options = ["All"] + sorted(df_rep["status"].dropna().astype(str).str.upper().unique().tolist())
            report_status = st.selectbox("Attendance status", status_options, key="report_status")
        if report_subject != "All":
            df_rep = df_rep[df_rep["subject"].astype(str) == report_subject]
        if report_section != "All":
            df_rep = df_rep[df_rep["section"].astype(str) == report_section]
        if report_status != "All":
            df_rep = df_rep[df_rep["status"].astype(str).str.upper() == report_status]
        
        # Color-code the status column
        def highlight_status(val):
            status = str(val).upper()
            if status == "PRESENT":
                return "background-color: #D6DDC9; color: #33402A"
            if status == "LATE":
                return "background-color: #E8D5AC; color: #624A20"
            if status == "ABSENT":
                return "background-color: #D7B6AB; color: #5A2923"
            if status == "INCOMPLETE":
                return "background-color: #BEB5A9; color: #291C0E"
            return ""

        styled = df_rep.style.map(highlight_status, subset=["status"])
        st.dataframe(styled, width="stretch")

        # Summary stats
        normalized_status = df_rep["status"].astype(str).str.upper()
        attended_count = int(normalized_status.isin(["PRESENT", "LATE"]).sum())
        eligible_count = int(normalized_status.isin(["PRESENT", "LATE", "ABSENT", "INCOMPLETE"]).sum())
        attendance_rate = (100.0 * attended_count / eligible_count) if eligible_count else 0.0
        period_count = len(df_rep[["section", "period"]].drop_duplicates()) if not df_rep.empty else 0
        col1, col2, col3, col4, col5, col6 = st.columns(6)
        with col1:
            st.metric("Periods", period_count)
        with col2:
            st.metric("Student-periods", eligible_count)
        with col3:
            present_count = int((normalized_status == "PRESENT").sum())
            st.metric("Present", present_count)
        with col4:
            late_count = int((normalized_status == "LATE").sum())
            st.metric("Late", late_count)
        with col5:
            absent_count = int(normalized_status.isin(["ABSENT", "INCOMPLETE"]).sum())
            st.metric("Absent / incomplete", absent_count)
        with col6:
            st.metric("Attendance rate", f"{attendance_rate:.1f}%")

        if not df_rep.empty:
            st.subheader("Period-wise summary")
            period_summary = (
                df_rep.assign(status=df_rep["status"].astype(str).str.upper())
                .pivot_table(
                    index=[column for column in ("section", "period", "subject", "faculty") if column in df_rep.columns],
                    columns="status", values="roll_number", aggfunc="count", fill_value=0,
                )
                .reset_index()
            )
            for status in ("PRESENT", "LATE", "ABSENT", "INCOMPLETE"):
                if status not in period_summary.columns:
                    period_summary[status] = 0
            period_summary["ELIGIBLE"] = period_summary[["PRESENT", "LATE", "ABSENT", "INCOMPLETE"]].sum(axis=1)
            period_summary["ATTENDANCE %"] = (
                100.0 * (period_summary["PRESENT"] + period_summary["LATE"])
                / period_summary["ELIGIBLE"].replace(0, np.nan)
            ).fillna(0).round(1)
            st.dataframe(period_summary, hide_index=True, width="stretch")

        st.markdown("---")
        st.subheader("Export report")
        col_ex1, col_ex2 = st.columns(2)

        with col_ex1:
            if EXCEL_AVAILABLE:
                excel_data = export_to_excel(df_rep, r_dept, date_str)
                downloaded = st.download_button(
                    label="Download Excel (.xlsx)",
                    data=excel_data,
                    file_name=f"{r_dept}_{date_str}_attendance.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                if downloaded:
                    record_system("REPORT EXPORTED", f"{r_dept}:{date_str}", {"department": r_dept, "date": date_str, "format": "xlsx"})
            else:
                st.warning("Install openpyxl for Excel export: `pip install openpyxl`")

        with col_ex2:
            if PDF_AVAILABLE:
                pdf_data = export_to_pdf(df_rep, r_dept, date_str)
                downloaded = st.download_button(
                    label="Download PDF",
                    data=pdf_data,
                    file_name=f"{r_dept}_{date_str}_attendance.pdf",
                    mime="application/pdf"
                )
                if downloaded:
                    record_system("REPORT EXPORTED", f"{r_dept}:{date_str}", {"department": r_dept, "date": date_str, "format": "pdf"})
            else:
                st.warning("Install reportlab for PDF export: `pip install reportlab`")

    else:
        st.info(f"No attendance recorded for **{r_dept}** on **{date_str}**.")


@st.cache_resource
def load_resources():
    recognizers    = {}
    student_lookup = {}
    trained_labels = {}

    for dept in DEPARTMENTS:
        for section in SECTIONS:
            section_key = f"{dept}_{section}"

            csv_path = BASE_DIR / dept / section / "students.csv"
            if csv_path.exists():
                try:
                    df = pd.read_csv(csv_path, dtype=str)
                    for _, r in df.iterrows():
                        roll = str(r.get("roll_number", "")).strip()
                        if roll:
                            student_lookup[(section_key, int(roll))] = {
                                "roll":    roll,
                                "name":    r.get("name", ""),
                                "dept":    dept,
                                "section": section,
                                "email":   r.get("email", "")
                            }
                            try:
                                pass  # Labels are scoped by department and section.
                            except Exception:
                                pass
                except Exception:
                    pass

            faces_dir = BASE_DIR / dept / section / "faces"
            local_labels = set()
            if faces_dir.exists():
                for s_dir in faces_dir.iterdir():
                    if s_dir.is_dir() and s_dir.name.isdigit() and any(s_dir.glob("*.jpg")):
                        local_labels.add(s_dir.name)
                        try:
                            local_labels.add(int(s_dir.name))
                        except Exception:
                            pass
            trained_labels[section_key] = local_labels

            model_path = BASE_DIR / dept / section / "trainer.yml"
            if model_path.exists() and local_labels:
                rec = create_lbph_recognizer()
                if rec:
                    rec.read(str(model_path))
                    recognizers[section_key] = rec

    return recognizers, student_lookup, trained_labels


def run_browser_attendance(video_placeholder, log_placeholder, class_context=None):
    """Recognize browser webcam frames without trying to open a server camera."""
    recognizers, student_lookup, trained_labels = load_resources()
    if not recognizers:
        st.error("No trained models found. Enroll students and train their section first.")
        st.session_state.attendance_active = False
        return

    scoped_key = f"{class_context.department}_{class_context.section}" if class_context else None
    if scoped_key and scoped_key not in recognizers:
        st.session_state.attendance_active = False
        st.error(
            f"No trained face model is available for {class_context.department} - "
            f"Section {class_context.section}. Train that section before starting recognition."
        )
        return

    face_cascade = get_face_detector()
    liveness = BlinkChallengeProvider()
    face_tracker = FaceTracker()
    last_recorded = {}
    config = attendance_rules()
    preview_only = bool(st.session_state.get("preview_only", False))
    processor_lock = threading.Lock()

    def process_browser_frame(video_frame):
        import av

        frame = video_frame.to_ndarray(format="bgr24")
        with processor_lock:
            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            faces = list(face_cascade.detectMultiScale(gray, 1.1, 8, minSize=(80, 80)))
            track_ids = face_tracker.assign(faces)
            liveness.expire(set(track_ids))

            for (x, y, w, h), track_id in zip(faces, track_ids):
                face_roi = gray[y:y+h, x:x+w]
                try:
                    live_result = liveness.update(track_id, face_roi)
                except (ValueError, cv2.error):
                    continue

                if config["liveness_required"] and live_result.status != LivenessStatus.LIVE:
                    if live_result.status == LivenessStatus.SPOOF:
                        color, label_text = (0, 0, 210), "SPOOF DETECTED"
                    else:
                        color = (0, 190, 230)
                        label_text = f"VERIFYING: {live_result.prompt}"
                    cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                    cv2.putText(frame, label_text, (x, max(25, y-10)), cv2.FONT_HERSHEY_SIMPLEX, .56, color, 2)
                    continue

                try:
                    normalized_face = prepare_face(face_roi)
                except (ValueError, cv2.error):
                    continue

                best_distance, best_match = 1000.0, None
                candidates = recognizers.items()
                if scoped_key:
                    candidates = [(scoped_key, recognizers[scoped_key])]
                for key, recognizer in candidates:
                    try:
                        label, distance = recognizer.predict(normalized_face)
                    except cv2.error:
                        continue
                    if distance < float(config["lbph_distance"]) and distance < best_distance:
                        labels = trained_labels.get(key, set())
                        if label in labels or str(label) in labels:
                            best_distance, best_match = float(distance), (key, label)

                color, label_text = (0, 0, 210), "UNKNOWN"
                info = student_lookup.get(best_match) if best_match else None
                if info:
                    color = (0, 180, 70)
                    label_text = f"LIVE | {info['name']} ({info['roll']}) | distance {best_distance:.1f}"
                    now_monotonic = time.monotonic()
                    if class_context and not preview_only:
                        if now_monotonic - last_recorded.get(info["roll"], 0) >= 5:
                            try:
                                record_observation(class_context, info["roll"], best_distance, True, base_dir=BASE_DIR)
                                last_recorded[info["roll"]] = now_monotonic
                            except (ValueError, OSError):
                                label_text = f"{info['name']} | ATTENDANCE WRITE FAILED"
                                color = (0, 120, 230)

                cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                cv2.putText(frame, label_text, (x, max(25, y-10)), cv2.FONT_HERSHEY_SIMPLEX, .56, color, 2)

        return av.VideoFrame.from_ndarray(frame, format="bgr24")

    # Third-party components calculate their iframe height from their direct
    # Streamlit block. Nesting WebRTC inside st.empty() can collapse the iframe
    # to an unclickable strip after a rerun, so clear the legacy image slot and
    # render the browser camera as a normal component block.
    video_placeholder.empty()
    context = browser_camera_component("smartattend-browser-attendance", process_browser_frame)
    if context is not None:
        if context.state.playing:
            log_placeholder.info("Browser camera connected. Keep this page open during attendance.")
        else:
            log_placeholder.info("Click START above and allow camera access in your browser.")


def run_attendance_loop(video_placeholder, log_placeholder, class_context=None):
    if browser_camera_enabled():
        return run_browser_attendance(video_placeholder, log_placeholder, class_context)

    recognizers, student_lookup, trained_labels = load_resources()

    if not recognizers:
        st.error("❌ No trained models found. Please go to Training page and train first.")
        st.session_state.attendance_active = False
        return

    cap          = open_camera()
    if cap is None:
        st.session_state.attendance_active = False
        return
    face_cascade = get_face_detector()

    liveness = BlinkChallengeProvider()
    face_tracker = FaceTracker()
    last_recorded = {}
    last_display_status = {}
    config = attendance_rules()
    scoped_key = f"{class_context.department}_{class_context.section}" if class_context else None
    if scoped_key and scoped_key not in recognizers:
        cap.release()
        st.session_state.attendance_active = False
        st.error(
            f"No trained face model is available for {class_context.department} - "
            f"Section {class_context.section}. Train that section before starting recognition."
        )
        return

    try:
        while st.session_state.attendance_active:
            if class_context and datetime.now() >= class_context.ends_at:
                finalize_due(datetime.now(), BASE_DIR)
                st.session_state.attendance_active = False
                log_placeholder.success(
                    f"{class_context.period} ended. Attendance was finalized and the report was updated."
                )
                break
            ret, frame = cap.read()
            if not ret:
                st.session_state.attendance_active = False
                st.error("Camera stopped delivering frames. Close other camera apps and restart recognition.")
                break
    
            gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            # Liveness must see closed-eye frames, so recognition uses the raw
            # conservative face detector rather than the enrollment eye filter.
            faces = list(face_cascade.detectMultiScale(gray, 1.1, 8, minSize=(80, 80)))
            track_ids = face_tracker.assign(faces)
            liveness.expire(set(track_ids))
    
            if not faces:
                log_placeholder.empty()

            for (x, y, w, h), track_id in zip(faces, track_ids):
                face_roi = gray[y:y+h, x:x+w]
                try:
                    live_result = liveness.update(track_id, face_roi)
                except Exception:
                    continue

                if config["liveness_required"] and live_result.status != LivenessStatus.LIVE:
                    if live_result.status == LivenessStatus.SPOOF:
                        color = (0, 0, 210)
                        name_text = "SPOOF DETECTED"
                    else:
                        color = (0, 190, 230)
                        name_text = f"VERIFYING: {live_result.prompt}"
                    cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                    cv2.putText(frame, name_text, (x, max(25, y-10)), cv2.FONT_HERSHEY_SIMPLEX, 0.56, color, 2)
                    continue

                try:
                    face_roi = prepare_face(face_roi)
                except Exception:
                    continue
    
                best_confidence = 1000
                best_match      = None
    
                candidates = recognizers.items()
                if scoped_key:
                    candidates = [(scoped_key, recognizers[scoped_key])] if scoped_key in recognizers else []
                for key, recognizer in candidates:
                    try:
                        label, confidence = recognizer.predict(face_roi)
                        if confidence < float(config["lbph_distance"]) and confidence < best_confidence:
                            section_labels = trained_labels.get(key, set())
                            if label in section_labels or str(label) in section_labels:
                                best_confidence = confidence
                                best_match      = (key, label)
                    except Exception:
                        pass
    
                color     = (0, 0, 210)
                name_text = "UNKNOWN"
    
                if best_match is not None:
                    info = student_lookup.get(best_match)
                    if info:
                        name_text = f"LIVE | {info['name']} ({info['roll']}) | distance {best_confidence:.1f}"
                        color     = (0, 255, 0)

                        now_monotonic = time.monotonic()
                        if class_context and not st.session_state.get("preview_only", False):
                            if now_monotonic - last_recorded.get(info["roll"], 0) >= 5:
                                status = record_observation(
                                    class_context,
                                    info["roll"],
                                    best_confidence,
                                    True,
                                    base_dir=BASE_DIR,
                                )
                                last_recorded[info["roll"]] = now_monotonic
                                if last_display_status.get(info["roll"]) != status:
                                    log_placeholder.info(
                                        f"{info['name']} · {class_context.subject} · {status.title()}"
                                    )
                                    last_display_status[info["roll"]] = status
    
                cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                cv2.putText(frame, name_text, (x, max(25, y-10)), cv2.FONT_HERSHEY_SIMPLEX, 0.56, color, 2)
    
            frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            video_placeholder.image(frame_rgb, channels="RGB")
    finally:
        cap.release()


def mark_attendance(info, period, period_start):
    date_str    = datetime.now().strftime("%Y-%m-%d")
    reports_dir = BASE_DIR / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    csv_path = reports_dir / f"{info['dept']}_{date_str}_attendance.csv"

    if csv_path.exists():
        df = pd.read_csv(csv_path, dtype={"roll_number": str})
    else:
        df = pd.DataFrame(columns=["date", "roll_number", "name", "section", "period", "sign_in_time", "status"])

    mask = (
        (df["date"] == date_str) &
        (df["roll_number"].astype(str) == str(info["roll"])) &
        (df["period"] == period)
    )

    if not mask.any():
        now        = datetime.now()
        sign_in_str = now.strftime("%Y-%m-%d %H:%M:%S")
        status     = get_attendance_status(sign_in_str, period_start)

        new_row = {
            "date":        date_str,
            "roll_number": info["roll"],
            "name":        info["name"],
            "section":     info["section"],
            "period":      period,
            "sign_in_time": sign_in_str,
            "status":      status
        }
        df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
        df.to_csv(csv_path, index=False)
        return status

    # Return existing status if already marked
    existing = df[mask]["status"].values
    return existing[0] if len(existing) > 0 else "Present"


# -----------------------
# Excel Export
# -----------------------
def export_to_excel(df, dept, date_str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    headers = ["Date", "Roll Number", "Name", "Section", "Semester", "Period"]
    if "subject" in df.columns:
        headers.extend(["Subject", "Subject Code", "Faculty"])
    headers.extend(["Class Start", "Class End", "Sign-In Time", "Last Seen", "Presence Duration", "Status"])

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"] = f"{dept} Department — Attendance Report — {date_str}"
    ws["A1"].font      = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row
    header_fill = PatternFill(start_color="2F75B6", end_color="2F75B6", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border    = border

    # Data rows
    present_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    late_fill    = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    absent_fill  = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    incomplete_fill = PatternFill(start_color="D9E2EC", end_color="D9E2EC", fill_type="solid")

    for row_idx, row in df.iterrows():
        status = str(row.get("status", "Present")).upper()
        fill = {"PRESENT": present_fill, "LATE": late_fill, "ABSENT": absent_fill,
                "INCOMPLETE": incomplete_fill}.get(status, incomplete_fill)

        for col_idx, header in enumerate(headers, start=1):
            col_key = header.lower().replace(" ", "_").replace("-", "_")
            # map header to df column
            col_map = {
                "date": "date", "roll_number": "roll_number", "name": "name",
                "section": "section", "semester": "semester", "period": "period", "subject": "subject",
                "subject_code": "subject_code", "faculty": "faculty", "class_start": "class_start",
                "class_end": "class_end", "sign_in_time": "sign_in_time", "last_seen": "last_seen",
                "presence_duration": "presence_duration", "status": "status"
            }
            val  = row.get(col_map.get(col_key, col_key), "")
            cell = ws.cell(row=row_idx + 3, column=col_idx, value=val)
            cell.fill      = fill
            cell.border    = border
            cell.alignment = Alignment(horizontal="center")

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# -----------------------
# PDF Export
# -----------------------
def export_to_pdf(df, dept, date_str):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1, fontSize=16)
    elements.append(Paragraph(f"{dept} Department — Attendance Report", title_style))
    elements.append(Paragraph(f"Date: {date_str}", styles["Normal"]))
    elements.append(Spacer(1, 0.2 * inch))

    # Summary
    normalized = df["status"].astype(str).str.upper() if "status" in df.columns else pd.Series(dtype=str)
    present_count = int((normalized == "PRESENT").sum())
    late_count = int((normalized == "LATE").sum())
    absent_count = int((normalized == "ABSENT").sum())
    incomplete_count = int((normalized == "INCOMPLETE").sum())
    elements.append(Paragraph(
        f"Total: {len(df)} | Present: {present_count} | Late: {late_count} | "
        f"Absent: {absent_count} | Incomplete: {incomplete_count}", styles["Normal"]
    ))
    elements.append(Spacer(1, 0.2 * inch))

    # Table
    if "subject" in df.columns:
        headers = ["Roll No", "Name", "Section", "Subject", "Period", "Sign-In Time", "Status"]
        col_map = ["roll_number", "name", "section", "subject", "period", "sign_in_time", "status"]
        col_widths = [55, 85, 42, 88, 48, 118, 62]
    else:
        headers = ["Roll No", "Name", "Section", "Period", "Sign-In Time", "Status"]
        col_map = ["roll_number", "name", "section", "period", "sign_in_time", "status"]
        col_widths = [70, 100, 60, 60, 120, 60]
    data    = [headers]

    for _, row in df.iterrows():
        data.append([str(row.get(c, "")) for c in col_map])

    t = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ("BACKGROUND",  (0, 0), (-1, 0),  colors.HexColor("#2F75B6")),
        ("TEXTCOLOR",   (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",    (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0, 0), (-1, 0),  10),
        ("ALIGN",       (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME",    (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",    (0, 1), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8F9FA"), colors.white]),
        ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("TOPPADDING",  (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]

    # Highlight Late rows in yellow
    for i, row in enumerate(df.itertuples(), start=1):
        if hasattr(row, "status") and str(row.status).upper() == "LATE":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF3CD")))
        elif hasattr(row, "status") and str(row.status).upper() == "ABSENT":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F8D7DA")))
        elif hasattr(row, "status") and str(row.status).upper() == "INCOMPLETE":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#D9E2EC")))

    t.setStyle(TableStyle(style_cmds))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)
    return buf.read()


# -----------------------
# Page: Email Alerts
# -----------------------
def page_email_alerts():
    ui.heading("COMMUNICATIONS", "Keep everyone informed.", "Manage attendance alerts and share daily summaries with your academic community.")

    st.info("""
    Configure your Gmail SMTP settings here (smtp.gmail.com, SSL port 465).
    Once configured, you can send absence alerts to students/parents.
    """)

    with st.form("email_config_form"):
        st.subheader("SMTP Settings")
        sender_email = st.text_input("Sender Email (Gmail)", placeholder="yourschool@gmail.com")
        app_password = st.text_input(
            "App Password",
            type="password",
            help="Use Gmail App Password (not your regular password). Enable 2FA first, then generate app password."
        )
        st.form_submit_button("Save settings (session only)")

    if sender_email:
        st.session_state["sender_email"] = sender_email
    if app_password:
        st.session_state["app_password"] = app_password.replace(" ", "")

    st.markdown("---")
    st.subheader("Absence alerts")

    col1, col2 = st.columns(2)
    with col1:
        alert_dept = st.selectbox("Department", DEPARTMENTS, key="alert_dept")
        alert_date = st.date_input("Date", datetime.now()).strftime("%Y-%m-%d")
    with col2:
        alert_report = refresh_report(datetime.strptime(alert_date, "%Y-%m-%d").date(), alert_dept, BASE_DIR)
        if alert_report and alert_report.exists():
            alert_frame = pd.read_csv(alert_report, dtype={"roll_number": str})
            period_options = sorted(alert_frame["period"].dropna().astype(str).unique().tolist())
        else:
            period_options = [f"Period{i}" for i in range(1, 13)]
        alert_period = st.selectbox("Period", period_options, key="alert_period")

    if st.button("Send absence alerts"):
        send_absence_alerts(alert_dept, alert_date, alert_period)

    st.markdown("---")
    st.subheader("Daily faculty summary")
    teacher_email = st.text_input("Teacher Email")
    summary_dept  = st.selectbox("Department for Summary", DEPARTMENTS, key="sum_dept")
    summary_date  = st.date_input("Summary Date", datetime.now(), key="sum_date").strftime("%Y-%m-%d")

    if st.button("Send summary email"):
        send_summary_email(teacher_email, summary_dept, summary_date)


def send_absence_alerts(dept, date_str, period):
    sender_email = st.session_state.get("sender_email", "")
    app_password = st.session_state.get("app_password", "")

    if not sender_email or not app_password:
        st.error("Please configure email settings first.")
        return

    # Prefer finalized explicit ABSENT rows from the attendance engine. Legacy
    # reports are still supported through the previous roster comparison.
    present_rolls = set()
    explicit_absent_rolls = None
    report_path = refresh_report(datetime.strptime(date_str, "%Y-%m-%d").date(), dept, BASE_DIR)

    if report_path and report_path.exists():
        df_rep = pd.read_csv(report_path, dtype={"roll_number": str})
        rows_in_period = df_rep[df_rep["period"] == period]
        normalized = rows_in_period["status"].astype(str).str.upper()
        present_rolls = set(
            rows_in_period.loc[normalized.isin(["PRESENT", "LATE"]), "roll_number"].astype(str)
        )
        if normalized.isin(["ABSENT", "INCOMPLETE"]).any():
            explicit_absent_rolls = set(
                rows_in_period.loc[normalized == "ABSENT", "roll_number"].astype(str)
            )

    # Get all students in dept
    all_students = []
    for section in SECTIONS:
        csv_path = BASE_DIR / dept / section / "students.csv"
        if csv_path.exists():
            df_s = pd.read_csv(csv_path, dtype=str)
            for _, r in df_s.iterrows():
                all_students.append({
                    "roll": str(r.get("roll_number", "")),
                    "name": r.get("name", ""),
                    "email": r.get("email", "")
                })

    if explicit_absent_rolls is None:
        absent_students = [s for s in all_students if s["roll"] not in present_rolls and s["email"]]
    else:
        absent_students = [s for s in all_students if s["roll"] in explicit_absent_rolls and s["email"]]

    if not absent_students:
        st.info("No absent students found (or no emails configured).")
        return

    sent = 0
    failed = 0
    for student in absent_students:
        try:
            msg            = MIMEMultipart()
            msg["From"]    = sender_email
            msg["To"]      = student["email"]
            msg["Subject"] = f"Attendance Alert — {date_str}"

            body = f"""
Dear Parent/Guardian,

This is to inform you that your ward {student['name']} (Roll: {student['roll']})
was ABSENT during {period} on {date_str} for {dept} department.

Please ensure regular attendance.

Regards,
College Attendance System
            """.strip()

            msg.attach(MIMEText(body, "plain"))
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20) as server:
                server.login(sender_email, app_password)
                server.sendmail(sender_email, student["email"], msg.as_string())
            sent += 1
        except Exception as e:
            failed += 1

    st.success(f"✅ Sent {sent} absence alerts. Failed: {failed}")


def send_summary_email(teacher_email, dept, date_str):
    sender_email = st.session_state.get("sender_email", "")
    app_password = st.session_state.get("app_password", "")

    if not sender_email or not app_password:
        st.error("Please configure email settings first.")
        return

    if not teacher_email:
        st.error("Please enter teacher email.")
        return

    report_path = refresh_report(datetime.strptime(date_str, "%Y-%m-%d").date(), dept, BASE_DIR)
    if not report_path or not report_path.exists():
        st.error("No attendance report found for selected date/department.")
        return

    df_rep = pd.read_csv(report_path, dtype={"roll_number": str})

    try:
        msg            = MIMEMultipart()
        msg["From"]    = sender_email
        msg["To"]      = teacher_email
        msg["Subject"] = f"Daily Attendance Summary — {dept} — {date_str}"

        normalized = df_rep["status"].astype(str).str.upper() if "status" in df_rep.columns else pd.Series(dtype=str)
        present_count = int((normalized == "PRESENT").sum())
        late_count = int((normalized == "LATE").sum())
        absent_count = int((normalized == "ABSENT").sum())
        incomplete_count = int((normalized == "INCOMPLETE").sum())
        total = len(df_rep)
        period_count = len(df_rep[["section", "period"]].drop_duplicates()) if total else 0
        period_lines = []
        if total:
            for period_name, period_rows in df_rep.groupby("period", sort=True):
                statuses = period_rows["status"].astype(str).str.upper()
                period_lines.append(
                    f"{period_name}: Present {(statuses == 'PRESENT').sum()}, "
                    f"Late {(statuses == 'LATE').sum()}, Absent {(statuses == 'ABSENT').sum()}, "
                    f"Incomplete {(statuses == 'INCOMPLETE').sum()}"
                )
        period_breakdown = "\n".join(period_lines) or "No completed periods"

        body = f"""
Dear Teacher,

Here is the attendance summary for {dept} Department on {date_str}:

Periods Held  : {period_count}
Student-Periods: {total}
Present       : {present_count}
Late          : {late_count}
Absent        : {absent_count}
Incomplete    : {incomplete_count}

Period-wise Summary:
{period_breakdown}

Please find the detailed CSV report attached.

Regards,
College Attendance System
        """.strip()

        msg.attach(MIMEText(body, "plain"))

        # Attach CSV
        with open(report_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f"attachment; filename={report_path.name}")
            msg.attach(part)

        with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=20) as server:
            server.login(sender_email, app_password)
            server.sendmail(sender_email, teacher_email, msg.as_string())

        st.success(f"✅ Summary email sent to {teacher_email}!")

    except Exception as e:
        st.error(f"Failed to send email: {e}")


# -----------------------
# Main App Structure
# -----------------------
def main():
    logo_path = Path(__file__).resolve().parent / "assets" / "smartattend_mark.png"
    page_icon = Image.open(logo_path) if logo_path.exists() else "◉"
    st.set_page_config(page_title="SmartAttend AI | Attendance", page_icon=page_icon, layout="wide")
    ui.apply_theme()
    ui.topbar()

    # Safe, idempotent startup work: preserve legacy CSVs while keeping SQLite
    # current and finalize any configured classes that ended while the app was open or closed.
    import_legacy(BASE_DIR)
    # A fresh cloud deployment has no private runtime database. Seed the
    # approved institutional AIML Semester V schedule once so timetable and
    # current-class resolution work immediately, while preserving any
    # timetable already configured locally or by an administrator.
    with connect() as db:
        timetable_count = db.execute("SELECT COUNT(*) FROM timetables").fetchone()[0]
    if timetable_count == 0:
        from scripts.import_aiml_sem5_timetable import import_timetable
        import_timetable()
    finalize_due(datetime.now(), BASE_DIR)

    # Check for missing optional libraries
    missing = []
    if not EXCEL_AVAILABLE:
        missing.append("`pip install openpyxl`")
    if not PDF_AVAILABLE:
        missing.append("`pip install reportlab`")
    if missing:
        st.sidebar.warning("⚠️ Optional libraries missing:\n" + "\n".join(missing))

    ui.brand()
    page = st.sidebar.radio(
        "WORKSPACE",
        ["📋 Attendance", "🎓 Training", "🗓 Timetable", "📈 Risk Analysis",
         "✏️ Corrections", "🛡 Audit Logs", "📧 Email Alerts", "🔐 Staff Access"],
        format_func=lambda p: {
            "📋 Attendance": "Live attendance",
            "🎓 Training": "Students & training",
            "🗓 Timetable": "Timetable & rules",
            "📈 Risk Analysis": "Risk analysis",
            "✏️ Corrections": "Attendance corrections",
            "🛡 Audit Logs": "Audit logs",
            "📧 Email Alerts": "Email alerts",
            "🔐 Staff Access": "Staff access",
        }[p],
    )
    st.sidebar.markdown("---")
    st.sidebar.caption("CAMERA SETTINGS")
    default_camera_source = 0 if os.name != "nt" else 1
    st.sidebar.selectbox(
        "Camera source",
        ["Browser camera", "Local OpenCV camera"],
        index=default_camera_source,
        key="camera_source",
        help="Use Browser camera on the deployed website and Local OpenCV camera when running on Windows.",
    )
    if not browser_camera_enabled():
        st.sidebar.number_input("Camera number", min_value=0, max_value=5, value=0, step=1, key="camera_index")
        st.sidebar.caption("0 is usually the built-in webcam. Stop capture before switching cameras.")
    else:
        st.sidebar.caption("Camera permission is requested by your browser when you start a video stream.")

    if page != "🎓 Training":
        st.session_state.training_active = False
    if page != "📋 Attendance":
        st.session_state.attendance_active = False

    if page == "🎓 Training":
        page_training()
    elif page == "🗓 Timetable":
        from smart_attendance.views.timetable_legacy import show
        show()
    elif page == "📈 Risk Analysis":
        from smart_attendance.views.risk_analysis import show
        show(st.session_state.get("staff_token"))
    elif page == "✏️ Corrections":
        from smart_attendance.views.corrections import show
        show(st.session_state.get("staff_token"))
    elif page == "🛡 Audit Logs":
        from smart_attendance.views.audit_logs import show
        show(st.session_state.get("staff_token"))
    elif page == "📧 Email Alerts":
        page_email_alerts()
    elif page == "🔐 Staff Access":
        from smart_attendance.views.staff_access import show
        show()
    else:
        page_attendance()

    st.sidebar.markdown("---")
    st.sidebar.markdown('<div class="sa-footer">SMARTATTEND AI<br>Face recognition · Academic operations</div>', unsafe_allow_html=True)


if __name__ == "__main__":
    main()
